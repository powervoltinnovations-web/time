import io
from html import escape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, A3, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from .domain import COLLEGE, Configuration
from .calendar import clock, minutes
from .resources import elective_entry, booking_key


def calendar_rows(settings):
    rows=[{'time':start+'–'+end,'slot':slot} for slot,(start,end) in enumerate(settings.periods())]
    rows += [{'time':pause.start+'–'+clock(minutes(pause.start)+pause.duration),'break':pause.name} for pause in settings.breaks]
    return sorted(rows,key=lambda row:row['time'])


def views(generation, view='group', target=None, batch=None):
    config=Configuration(**generation['config'])
    data=config.model_dump()
    settings=config.settings
    entries=generation['result']['entries']
    if batch:
        selected_batch=next((b for b in config.batches if b.id==batch),None)
        if view!='group' or not selected_batch or selected_batch.group_id!=target:
            raise ValueError('Choose a batch belonging to the selected division.')
        entries=[e for e in entries if not e.get('batch_id') or e['batch_id']==batch]
    scheduled={e['group_id'] for e in entries}
    if view not in ('group','teacher','classroom','lab'):
        raise ValueError('Choose a division, professor, classroom, or lab view.')
    room_view=view in ('classroom','lab')
    key='rooms' if room_view else 'groups' if view=='group' else 'teachers'
    records=[item for item in data[key] if (not target or item['id']==target) and (view!='group' or item['id'] in scheduled) and (not room_view or item['kind']==view)]
    if not records:
        raise ValueError('The selected division, professor, classroom, or lab is not in this saved timetable.')
    maps={name:{x['id']:x for x in data[name]} for name in ('teachers','subjects','groups','rooms','departments','batches')}
    output=[]
    for record in records:
        title=(f"{maps['departments'][record['department_id']]['code']} · Semester {record['semester']} · Division {record['division']}"
               if view=='group' else record['name'])
        if batch:
            title+=' · Batch '+selected_batch.name
        selected=[e for e in entries if e['room_id' if room_view else view+'_id']==record['id']]
        if view!='group':
            seen=set();unique=[];offerings={o.id:o for o in config.offerings}
            for e in selected:
                identity=(booking_key(config,offerings[e['offering_id']],e['kind']),e['day'],e['slot'],e['duration'],e['room_id'])
                if identity not in seen: unique.append(e);seen.add(identity)
            selected=unique
        grid=[['Time',*settings.days]]
        for row in calendar_rows(settings):
            if 'break' in row:
                grid.append([row['time'], *[row['break'].upper()]*len(settings.days)])
                continue
            slot=row['slot']
            cells=[row['time']]
            for day in range(len(settings.days)):
                blocked=next((b for b in settings.blocked if b.day==settings.days[day] and b.slot==slot),None)
                if blocked:
                    cells.append('BLOCKED\n'+blocked.reason)
                    continue
                matches=[e for e in selected if e['day']==day and e['slot']<=slot<e['slot']+e['duration']]
                texts=[]
                for e in matches:
                    subject,group=maps['subjects'][e['subject_id']],maps['groups'][e['group_id']]
                    kind='COMBINED A+B ELECTIVE' if elective_entry(config,e) else e['kind'].upper()+(' · ELECTIVE' if subject['elective'] else '')
                    who=maps['teachers'].get(e['teacher_id'],{}).get('name','') if view=='group' else f"{maps['departments'][group['department_id']]['code']} · Sem {group['semester']} · Div {group['division']}"
                    if room_view:
                        who+='\n'+maps['teachers'].get(e['teacher_id'],{}).get('name','')
                    batch_label=' · '+maps['batches'][e['batch_id']]['name'] if e.get('batch_id') else ''
                    texts.append(f"{subject['code']} · {kind}\n{subject['name']}\n{who}{batch_label}\n{maps['rooms'].get(e['room_id'],{}).get('name','')}")
                cells.append('\n—\n'.join(texts))
            grid.append(cells)
        output.append((title,grid,sum(e['duration']*settings.period_minutes/60 for e in selected)))
    return output


def workflow_label(generation):
    workflow=generation.get('workflow',{'state':'draft','locked':False})
    return workflow['state'].upper()+(' · LOCKED' if workflow['locked'] else '')


def excel(generation,view='group',target=None,batch=None):
    workbook=Workbook()
    workbook.remove(workbook.active)
    for index,(title,grid,hours) in enumerate(views(generation,view,target,batch),1):
        sheet=workbook.create_sheet(f'{index} '+{'group':'Division','teacher':'Professor','classroom':'Classroom','lab':'Lab'}[view])
        columns=len(grid[0])
        last=get_column_letter(columns)
        for row,text in [(1,COLLEGE),(2,f'UniSched | {title} | {workflow_label(generation)}'),(3,f"{generation['name']} | {hours:g} scheduled teaching hours / week")]:
            sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=columns)
            cell=sheet.cell(row,1,text)
            cell.data_type='s'
            cell.font=Font(name='Calibri',size=14 if row==1 else 11,bold=row!=3,color='252645')
            cell.alignment=Alignment(horizontal='center',vertical='center')
            sheet.row_dimensions[row].height=25
        settings=Configuration(**generation['config']).settings
        break_names={b.name.upper() for b in settings.breaks}
        for r,values in enumerate(grid,5):
            pause=all(value in break_names for value in values[1:])
            for c,value in enumerate(values,1):
                cell=sheet.cell(r,c,value)
                cell.data_type='s'
                cell.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
                cell.font=Font(name='Calibri',size=10,color='FFFFFF' if r==5 else '252645',bold=r==5 or c==1)
                cell.fill=PatternFill('solid',fgColor='5046D6' if r==5 else 'F0EEFF' if pause else 'F6F7FB')
                cell.border=Border(bottom=Side(style='thin',color='DFE2EB'))
            sheet.row_dimensions[r].height=25 if r==5 or pause else max(78,max(value.count('\n') for value in values)*13+20)
        sheet.column_dimensions['A'].width=19
        for c in range(2,columns+1):
            sheet.column_dimensions[get_column_letter(c)].width=32
        sheet.freeze_panes='B6'
        sheet.sheet_view.showGridLines=False
        sheet.page_setup.orientation='landscape'
        sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth=1
        sheet.page_setup.fitToHeight=1
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.print_options.horizontalCentered=True
        sheet.print_area=f'A1:{last}{len(grid)+4}'
        sheet.oddFooter.center.text='UniSched · &P / &N'
    output=io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def pdf(generation,view='group',target=None,batch=None):
    output=io.BytesIO()
    large=view=='group' and bool(generation['config'].get('batches')) and not batch
    document=SimpleDocTemplate(output,pagesize=landscape(A3 if large else A4),leftMargin=24,rightMargin=24,topMargin=20,bottomMargin=20,title='UniSched timetable',author=COLLEGE)
    heading=ParagraphStyle('heading',fontName='Helvetica-Bold',fontSize=13,leading=17,textColor=colors.HexColor('#252645'),alignment=TA_CENTER)
    subheading=ParagraphStyle('subheading',fontName='Helvetica',fontSize=9,leading=13,alignment=TA_CENTER)
    cellstyle=ParagraphStyle('cell',fontName='Helvetica',fontSize=6.2 if large else 6.8,leading=6.8 if large else 8.2,alignment=TA_CENTER)
    headerstyle=ParagraphStyle('header',parent=cellstyle,fontName='Helvetica-Bold',textColor=colors.white)
    story=[]
    settings=Configuration(**generation['config']).settings
    break_names={b.name.upper() for b in settings.breaks}
    for index,(title,grid,hours) in enumerate(views(generation,view,target,batch)):
        if index:
            story.append(PageBreak())
        story += [Paragraph(escape(COLLEGE),heading),Spacer(1,5),Paragraph(escape(f"UniSched | {title} | {generation['name']} | {workflow_label(generation)}"),subheading),
                  Paragraph(escape(f"{hours:g} scheduled teaching hours/week · {settings.start}–{settings.end}"),subheading),Spacer(1,10)]
        table=Table([[Paragraph(escape(value).replace('\n','<br/>'),headerstyle if r==0 else cellstyle) for value in row] for r,row in enumerate(grid)],
                    colWidths=[74]+[(document.width-74)/(len(grid[0])-1)]*(len(grid[0])-1),repeatRows=1)
        commands=[('BACKGROUND',(0,0),(-1,0),colors.HexColor('#5046D6')),
                  ('BACKGROUND',(0,1),(-1,-1),colors.HexColor('#F7F8FC')),
                  ('VALIGN',(0,0),(-1,-1),'MIDDLE'),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#DDE0E9')),
                  ('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),
                  ('TOPPADDING',(0,0),(-1,-1),2 if large else 7),('BOTTOMPADDING',(0,0),(-1,-1),2 if large else 7)]
        for r,row in enumerate(grid[1:],1):
            if all(value in break_names for value in row[1:]):
                commands.append(('BACKGROUND',(0,r),(-1,r),colors.HexColor('#EDEBFF')))
        table.setStyle(TableStyle(commands))
        story += [table,Spacer(1,7),Paragraph('Empty cells indicate no scheduled teaching. Multi-period sessions repeat across their occupied periods. Batch sessions are identified by name.',subheading)]
    def footer(canvas,doc):
        canvas.saveState()
        canvas.setFont('Helvetica',6)
        canvas.drawCentredString(doc.pagesize[0]/2,10,COLLEGE+' | '+workflow_label(generation))
        canvas.restoreState()
    document.build(story,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue()
