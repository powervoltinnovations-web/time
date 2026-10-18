"""Full-fidelity Excel/JSON exchange and extended, editable assignment CSVs."""
import csv
import io
import json
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
from .domain import Configuration, CSV_FIELDS, parse_csv

EXTRA_FIELDS = ['department_code','department_name','student_strength','teacher_daily_limit',
                'teacher_email','teacher_short_code','teacher_lab_eligible','project_sessions_per_week',
                'batch_name','batch_strength','elective_set','course_type','weekly_hours',
                'classroom_pool','lab_pool','calendar_json','department_pool','teacher_department_code','elective_set_code','project_hours_per_session','project_room']
V2_FIELDS = CSV_FIELDS + EXTRA_FIELDS
TABLES = ['departments','teachers','subjects','rooms','groups','batches','elective_sets','offerings']


def safe(value):
    return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@','\t','\r')) else value


def configuration_csv_v2(config):
    m = {key:{x.id:x for x in getattr(config,key)} for key in TABLES}
    output = io.StringIO()
    writer = csv.DictWriter(output,fieldnames=V2_FIELDS)
    writer.writeheader()
    # The pools contain structured room metadata so automatic allocation survives import.
    pools = {kind:json.dumps([r.model_dump() for r in config.rooms if r.kind == kind],separators=(',',':')) for kind in ('classroom','lab')}
    for index,o in enumerate(config.offerings):
        g,s,t = m['groups'][o.group_id],m['subjects'][o.subject_id],m['teachers'][o.teacher_id]
        dept = m['departments'][g.department_id]
        row = dict(zip(CSV_FIELDS,[g.semester,g.division,t.name,t.max_hours,t.arrival,t.departure,s.name,s.code,
            str(s.elective).lower(),m['rooms'][g.room_id].name if g.room_id else 'AUTO',
            m['rooms'][o.lab_room_id].name if o.lab_room_id else '',o.lectures,o.labs]))
        row.update(department_code=dept.code,department_name=dept.name,student_strength=g.strength,
            teacher_daily_limit=t.max_daily_slots,teacher_email=t.email,teacher_short_code=t.short_code,
            teacher_lab_eligible=str(t.lab_eligible).lower(),project_sessions_per_week=o.projects,project_hours_per_session=o.project_hours,project_room=m['rooms'][o.project_room_id].name if o.project_room_id else '',
            batch_name=m['batches'][o.batch_id].name if o.batch_id else '',
            batch_strength=m['batches'][o.batch_id].strength if o.batch_id else '',
            elective_set=m['elective_sets'][s.elective_set_id].name if s.elective_set_id else '',
            course_type=s.course_type,weekly_hours=s.weekly_hours,
            classroom_pool=pools['classroom'] if index == 0 else '',lab_pool=pools['lab'] if index == 0 else '',
            calendar_json=config.settings.model_dump_json() if index == 0 else '',
            department_pool=json.dumps([d.model_dump() for d in config.departments]) if index == 0 else '',
            teacher_department_code=m['departments'][t.department_id].code,
            elective_set_code=m['elective_sets'][s.elective_set_id].code if s.elective_set_id else '')
        writer.writerow({k:safe(v) for k,v in row.items()})
    return output.getvalue()


def parse_extended_csv(content):
    rows = list(csv.DictReader(io.StringIO(content)))
    if not rows or len(rows)>240:
        raise ValueError('Provide 1–240 assignment rows.')
    if list(rows[0]) not in (V2_FIELDS, V2_FIELDS[:-1], V2_FIELDS[:-2]):
        raise ValueError('CSV headers must match the original or updated sample template exactly.')
    data = {key:[] for key in TABLES}
    data['settings'] = {}
    pools = {}
    departments = {}
    for row in rows:
        if None in row or any(value is None for value in row.values()):
            raise ValueError('Incorrect number of CSV columns.')
        if row['department_pool'].strip():
            for department in json.loads(row['department_pool']):
                old = departments.get(department['id'])
                if old and old != department:
                    raise ValueError('Inconsistent department pool.')
                departments[department['id']] = department
        for key in ('classroom_pool','lab_pool'):
            if row[key].strip():
                entries = json.loads(row[key])
                if not isinstance(entries,list):
                    raise ValueError('Room pools must be JSON lists from the template.')
                for room in entries:
                    if room['id'] in pools and pools[room['id']] != room:
                        raise ValueError('Inconsistent repeated room pool data.')
                    pools[room['id']] = room
        if row['calendar_json'].strip():
            calendar = json.loads(row['calendar_json'])
            if data['settings'] and data['settings'] != calendar:
                raise ValueError('Inconsistent calendar settings.')
            data['settings'] = calendar
    # Preserve department IDs so room and faculty eligibility survives CSV exchange.
    data['departments'] = list(departments.values())
    data['rooms'] = list(pools.values())
    indexes = {key:{} for key in TABLES}
    for department in data['departments']:
        indexes['departments'][department['code'].casefold()] = department
    for room in data['rooms']:
        indexes['rooms'][room['name'].casefold()] = room
    def add(table,key,values,identifier=None):
        old = indexes[table].get(key)
        if old:
            if {k:v for k,v in old.items() if k!='id'} != values:
                raise ValueError(f'Conflicting repeated {table} details for {key}.')
            return old['id']
        if identifier is None:
            index = len(data[table])+1
            existing_ids = {item['id'] for item in data[table]}
            while table[0]+str(index) in existing_ids:
                index += 1
            identifier = table[0]+str(index)
        item = {'id':identifier,**values}
        indexes[table][key] = item
        data[table].append(item)
        return item['id']
    def boolean(value):
        if value.lower() not in ('true','false','yes','no','1','0'):
            raise ValueError('Boolean columns must contain true/false, yes/no, or 1/0.')
        return value.lower() in ('true','yes','1')
    for number,row in enumerate(rows,2):
        try:
            row = {k:v.strip() for k,v in row.items()}
            code = row['department_code'] or 'GEN'
            dep_id = code.lower().replace(' ','_')
            dept = add('departments',code.casefold(),{'code':code,'name':row['department_name'] or code},dep_id)
            teacher_department = indexes['departments'].get((row['teacher_department_code'] or code).casefold())
            if teacher_department is None:
                raise ValueError('Teacher department must be declared in the department pool.')
            sem = int(row['semester'])
            teacher = add('teachers',row['teacher_name'].casefold(),{
                'name':row['teacher_name'],'max_hours':int(row['teacher_hours_per_week']),
                'arrival':row['teacher_arrival'],'departure':row['teacher_departure'],'department_id':teacher_department['id'],
                'max_daily_slots':int(row['teacher_daily_limit'] or 8),'email':row['teacher_email'],
                'short_code':row['teacher_short_code'],'lab_eligible':boolean(row['teacher_lab_eligible'] or 'true')})
            elective_set = None
            if row['elective_set']:
                elective_set = add('elective_sets',(dept,sem,row['elective_set']),{'department_id':dept,'semester':sem,
                    'name':row['elective_set'],'code':row['elective_set_code'] or row['elective_set'][:30]})
            subject = add('subjects',(dept,row['subject_code'].casefold()),{'name':row['subject_name'],
                'code':row['subject_code'],'semester':sem,'elective':boolean(row['elective']),'department_id':dept,
                'elective_set_id':elective_set,'course_type':row['course_type'] or 'theory','weekly_hours':int(row['weekly_hours'] or 3)})
            def room_id(name,kind):
                if name.upper() in ('','AUTO'):
                    return None
                existing = indexes['rooms'].get(name.casefold())
                if existing:
                    if kind != 'any' and existing['kind'] != kind:
                        raise ValueError('Room type does not match its use.')
                    return existing['id']
                return add('rooms',name.casefold(),{'name':name,'kind':'classroom' if kind=='any' else kind,'capacity':80,'department_id':None,'eligible_years':[]})
            classroom,lab = room_id(row['classroom'],'classroom'),room_id(row['lab_room'],'lab')
            group = add('groups',(dept,sem,row['division'].casefold()),{'department_id':dept,'semester':sem,
                'division':row['division'],'strength':int(row['student_strength'] or 60),'room_id':classroom})
            batch = None
            if row['batch_name']:
                batch = add('batches',(group,row['batch_name']),{'group_id':group,'name':row['batch_name'],'strength':int(row['batch_strength'])})
            data['offerings'].append({'id':f'o{len(data["offerings"])+1}','teacher_id':teacher,'subject_id':subject,
                'group_id':group,'lectures':int(row['lectures_per_week']),'labs':int(row['labs_per_week']),
                'projects':int(row['project_sessions_per_week'] or 0),
                'project_room_id':room_id(row.get('project_room',''),'any'),
                'project_hours':float(row['project_hours_per_session']) if row.get('project_hours_per_session') else None,'batch_id':batch,'lab_room_id':lab})
        except (KeyError,ValueError,TypeError) as exc:
            raise ValueError(f'CSV row {number}: {exc}') from exc
    return Configuration(**data)


def configuration_excel(config):
    book = Workbook()
    book.remove(book.active)
    for key in TABLES:
        sheet = book.create_sheet(key)
        rows = [item.model_dump() for item in getattr(config,key)]
        annotation = Configuration.model_fields[key].annotation.__args__[0]
        headers = list(annotation.model_fields)
        sheet.append(headers)
        for row in rows:
            sheet.append([json.dumps(row[h]) if isinstance(row[h],(dict,list)) else row[h] for h in headers])
        sheet.freeze_panes = 'A2'
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type = 's'
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
    sheet = book.create_sheet('settings')
    sheet.append(['setting','value_json'])
    for key,value in config.settings.model_dump().items():
        sheet.append([key,json.dumps(value)])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def parse_upload(raw, filename):
    if len(raw)>1_000_000:
        raise ValueError('The import file must be smaller than 1 MB.')
    if filename.lower().endswith('.json'):
        value = json.loads(raw)
        if not isinstance(value,dict):
            raise ValueError('JSON must contain an academic configuration object.')
        return Configuration(**value.get('configuration',value))
    if filename.lower().endswith('.xlsx'):
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(f.file_size for f in archive.infolist())>8_000_000:
                raise ValueError('The expanded Excel workbook is too large.')
        book = load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
        try:
            if not set(TABLES+['settings']) <= set(book.sheetnames):
                raise ValueError('Use the academic Excel template exported by UniSched, with all named worksheets.')
            data = {}
            for key in TABLES+['settings']:
                rows = book[key].iter_rows()
                header = [c.value for c in next(rows)]
                if len(header)>30:
                    raise ValueError('Too many worksheet columns.')
                data[key] = {} if key=='settings' else []
                for index,cells in enumerate(rows):
                    if index>=500 or any(c.data_type=='f' for c in cells):
                        raise ValueError('Use at most 500 rows per sheet and literal values instead of formulas.')
                    values = [c.value for c in cells]
                    if all(v is None for v in values):
                        continue
                    row = dict(zip(header,values))
                    if key=='settings':
                        data[key][row['setting']] = json.loads(row['value_json'])
                    else:
                        for field in ('eligible_years',):
                            if field in row:
                                row[field] = json.loads(row[field] or '[]')
                        for field in ('email','short_code'):
                            if field in row and row[field] is None:
                                row[field] = ''
                        data[key].append(row)
            return Configuration(**data)
        finally:
            book.close()
    return parse_csv(raw)
