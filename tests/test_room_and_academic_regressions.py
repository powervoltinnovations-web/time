"""Regressions using only the academic data shipped in this ZIP."""
from copy import deepcopy
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app import storage
from app.analytics import analyze
from app.demo import presentation_config
from app.domain import Configuration
from app.main import app
from app.resources import eligible_rooms, booking_key, session_specs
from app.scheduler import generate, preflight


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("UNISCHED_DB", str(tmp_path / "test.db"))
    with TestClient(app) as client:
        yield client


@pytest.fixture(scope="module")
def valid_plan():
    config = presentation_config()
    config.settings.classroom_run_length = 3
    group = next(g for g in config.groups if g.semester == 3 and g.division == "A")
    result = generate(config, 30, [group.id], random_seed=7)
    assert result["status"] == "feasible", result["issues"]
    return {"config": config.model_dump(), "result": result}


def test_legacy_classroom_practical_flags_cannot_change_room_types():
    config = presentation_config()
    config.settings.practical_classroom_codes = [s.code for s in config.subjects]
    for offering in config.offerings:
        for kind, count, duration in session_specs(config, offering):
            if count and kind in ("lecture", "lab", "tutorial"):
                rooms = eligible_rooms(config, offering, kind)
                assert rooms
                assert {r.kind for r in rooms} == ({"lab"} if kind == "lab" else {"classroom"})
                if kind == "lab":
                    assert duration * config.settings.period_minutes == 120
                    assert booking_key(config, offering, kind) == ("single", offering.id, kind)


def test_full_zip_capacity_conflicts_are_reported_without_relaxation(client):
    before = client.get("/api/config").json()
    response = client.post("/api/generations", json={"flexible_practicals": True})
    assert response.status_code == 200
    record = response.json()
    assert record["result"]["status"] == "infeasible"
    assert record["result"]["entries"] == []
    assert record["config"]["settings"]["practical_classroom_codes"] == []
    assert record["config"]["settings"]["lab_slots"] == 2
    issues = record["result"]["issues"]
    room = next(i for i in issues if i["code"] == "ROOM_POOL_CAPACITY")
    assert "246 required room-hours" in room["problem"]
    assert "240 available room-hours" in room["problem"]
    assert "6 laboratories" in room["problem"]
    assert any("Pallavi M. N." in i["problem"] and "34 periods" in i["problem"] for i in issues)
    assert client.get("/api/config").json() == before


def test_generated_zip_division_preserves_counts_room_types_and_continuity(valid_plan):
    audit = analyze(valid_plan)
    assert audit["hard_conflicts"] == 0, audit["violations"]
    assert audit["classroom_optimization"]["early_changes"] == 0
    assert audit["classroom_optimization"]["preferred_periods"] == 3
    rooms = {r["id"]: r["kind"] for r in valid_plan["config"]["rooms"]}
    entries = valid_plan["result"]["entries"]
    assert len(entries) == 38
    assert any(e["kind"] == "lab" for e in entries)
    for entry in entries:
        if entry["kind"] == "lab":
            assert entry["duration"] == 2
            assert rooms[entry["room_id"]] == "lab"
        elif entry["kind"] in ("lecture", "tutorial"):
            assert rooms[entry["room_id"]] == "classroom"


@pytest.mark.parametrize("semester,sessions", [(3, 114), (5, 111), (7, 64)])
def test_each_semester_jointly_uses_shared_resources(semester, sessions):
    # These are separate scopes, not a claim that all eight divisions fit together.
    config = presentation_config()
    config.settings.classroom_run_length = 3
    result = generate(config, 30, [g.id for g in config.groups if g.semester == semester],
                      random_seed=7)
    if semester in (3, 5):
        # Reviewed reassignments exceed unchanged faculty caps in these scopes.
        # Report those conflicts; do not silently increase the faculty limits.
        assert result["status"] == "infeasible"
        assert result["entries"] == []
        assert any(i["code"] == "TEACHER_LOAD" for i in result["issues"])
        return
    assert result["status"] == "feasible", result["issues"]
    assert len(result["entries"]) == sessions
    audit = analyze({"config": config.model_dump(), "result": result})
    assert audit["hard_conflicts"] == 0, audit["violations"]
    assert audit["classroom_optimization"]["early_changes"] == 0


@pytest.mark.parametrize("kind,wrong_type", [("lab", "classroom"), ("lecture", "lab")])
def test_audit_and_manual_edit_reject_swapped_room_types(client, valid_plan, kind, wrong_type):
    record = deepcopy(valid_plan)
    record["config"]["settings"]["practical_classroom_codes"] = [
        s["code"] for s in record["config"]["subjects"]
    ]
    idx = next(i for i, e in enumerate(record["result"]["entries"]) if e["kind"] == kind)
    entry = record["result"]["entries"][idx]
    wrong_room = next(r["id"] for r in record["config"]["rooms"] if r["kind"] == wrong_type)
    original_room = entry["room_id"]
    entry["room_id"] = wrong_room
    assert any("ineligible room" in v for v in analyze(record)["violations"])
    entry["room_id"] = original_room
    identifier = storage.save_generation("Room validation", 1, record["config"], record["result"])
    response = client.post(f"/api/generations/{identifier}/edits", json={
        "revision": 0,
        "changes": [{"entry_index": idx, "day": entry["day"], "slot": entry["slot"],
                     "room_id": wrong_room}],
    })
    assert response.status_code == 422
    assert "ineligible room" in response.json()["detail"]
    assert len(client.get("/api/generations").json()) == 1
    assert client.get(f"/api/generations/{identifier}").json()["result"] == record["result"]


def test_invalid_saved_lab_placements_are_not_reused(valid_plan):
    config = Configuration(**valid_plan["config"])
    entries = deepcopy(valid_plan["result"]["entries"])
    lab = next(e for e in entries if e["kind"] == "lab")
    lab["room_id"] = next(r.id for r in config.rooms if r.kind == "classroom")
    result = generate(config, 30, valid_plan["result"]["scope"]["target_group_ids"],
                      initial_entries=entries, random_seed=7)
    assert result["status"] == "feasible", result["issues"]
    assert not result["summary"]["reused_placements"]
    assert analyze({"config": config.model_dump(), "result": result})["hard_conflicts"] == 0


def test_final_api_audit_never_saves_a_bad_candidate(client, valid_plan, monkeypatch):
    result = deepcopy(valid_plan["result"])
    next(e for e in result["entries"] if e["kind"] == "lab")["room_id"] = next(
        r["id"] for r in valid_plan["config"]["rooms"] if r["kind"] == "classroom"
    )
    monkeypatch.setattr("app.main.generate", lambda *args, **kwargs: result)
    response = client.post("/api/generations", json={})
    assert response.status_code == 500
    assert "failed final validation" in response.json()["detail"]
    assert client.get("/api/generations").json() == []


def test_subject_edits_deletion_revision_and_snapshots(client, valid_plan):
    current = client.get("/api/config").json()
    identifier = storage.save_generation("Snapshot", current["revision"],
                                         valid_plan["config"], valid_plan["result"])
    snapshot = client.get(f"/api/generations/{identifier}").json()
    subject = current["configuration"]["subjects"][0]
    subject["name"] = "Edited subject name"
    updated = client.put("/api/config", json=current)
    assert updated.status_code == 200
    assert client.get("/api/config").json()["configuration"]["subjects"][0]["name"] == subject["name"]
    assert client.put("/api/config", json=current).status_code == 409
    current = updated.json()
    next_config = deepcopy(current)
    next_config["configuration"]["subjects"] = [
        s for s in next_config["configuration"]["subjects"] if s["id"] != subject["id"]
    ]
    assert client.put("/api/config", json=next_config).status_code == 422
    next_config["configuration"]["offerings"] = [
        o for o in next_config["configuration"]["offerings"] if o["subject_id"] != subject["id"]
    ]
    assert client.put("/api/config", json=next_config).status_code == 200
    stored = client.get("/api/config").json()["configuration"]
    assert stored["offerings"] == next_config["configuration"]["offerings"]
    assert client.get(f"/api/generations/{identifier}").json() == snapshot


def test_timetable_exports_are_readable(client, valid_plan):
    identifier = storage.save_generation("Strict rooms", 1, valid_plan["config"], valid_plan["result"])
    pdf = client.get(f"/api/generations/{identifier}/export/pdf")
    assert pdf.status_code == 200
    pages = PdfReader(BytesIO(pdf.content)).pages
    assert pages and "DATTA MEGHE" in pages[0].extract_text().upper()
    excel = client.get(f"/api/generations/{identifier}/export/xlsx")
    assert excel.status_code == 200
    assert load_workbook(BytesIO(excel.content)).sheetnames
