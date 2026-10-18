from copy import deepcopy
from io import BytesIO
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from app import storage
from app.exports import views
from app.main import app
from test_scheduler import small_config


@pytest.fixture
def saved(tmp_path,monkeypatch):
    monkeypatch.setenv('UNISCHED_DB',str(tmp_path/'rooms.db'))
    config=small_config()
    config.rooms.extend([config.rooms[0].model_copy(update={'id':'project_room','name':'Project classroom'}),
                         config.rooms[0].model_copy(update={'id':'empty','name':'Empty classroom'})])
    config.offerings[0].projects=1
    config.offerings[0].project_hours=2
    config.offerings[0].project_room_id='project_room'
    config.teachers.append(config.teachers[0].model_copy(update={'id':'t2','name':'Second Teacher'}))
    config.groups.append(config.groups[0].model_copy(update={'id':'g2','division':'B'}))
    config.offerings.append(config.offerings[0].model_copy(update={'id':'o2','group_id':'g2','teacher_id':'t2','projects':0}))
    storage.initialize();storage.save_config(config.model_dump(),0)
    with TestClient(app) as client:
        response=client.post('/api/generations',json={'time_limit':20})
        assert response.status_code==200,response.text
        record=response.json();assert record['result']['status']=='feasible',record
        yield client,record


@pytest.mark.parametrize('view,room',[('classroom','r'),('classroom','project_room'),('lab','l'),('classroom','empty')])
def test_room_grids_contain_only_that_rooms_sessions(saved,view,room):
    client,record=saved;before=deepcopy(record)
    title,grid,hours=views(record,view,room)[0]
    entries=[e for e in record['result']['entries'] if e['room_id']==room]
    assert hours==sum(e['duration'] for e in entries)
    cells=[cell for row in grid[1:] for cell in row[1:] if 'Algorithms' in cell]
    assert len(cells)==sum(e['duration'] for e in entries)
    for entry in entries:
        teacher=next(t for t in record['config']['teachers'] if t['id']==entry['teacher_id'])
        group=next(g for g in record['config']['groups'] if g['id']==entry['group_id'])
        assert any(teacher['name'] in cell and f"Div {group['division']}" in cell for cell in cells)
    if room=='project_room': assert all('PROJECT' in cell for cell in cells)
    assert record==before
    assert client.get(f"/api/generations/{record['id']}").json()==before


@pytest.mark.parametrize('view,room',[('classroom','r'),('lab','l')])
def test_room_pdf_and_excel_exports(saved,view,room):
    client,record=saved
    for format in ('pdf','xlsx'):
        response=client.get(f"/api/generations/{record['id']}/export/{format}?view={view}&target={room}")
        assert response.status_code==200,response.text
        if format=='pdf':
            pages=PdfReader(BytesIO(response.content)).pages
            assert len(pages)==1
            text=pages[0].extract_text()
        else:
            book=load_workbook(BytesIO(response.content))
            assert book.sheetnames==['1 '+view.capitalize()]
            text=' '.join(str(c.value) for row in book.active for c in row if c.value)
        assert 'Professor Test' in text and 'Second Teacher' in text
        assert 'Div A' in text and 'Div B' in text
    all_rooms=client.get(f"/api/generations/{record['id']}/export/xlsx?view={view}")
    assert all_rooms.status_code==200
    book=load_workbook(BytesIO(all_rooms.content))
    assert len(book.worksheets)==(3 if view=='classroom' else 1)


@pytest.mark.parametrize('query',['view=lab&target=r','view=classroom&target=l','view=lab&target=missing',
                                'view=classroom&target=r&batch=missing'])
def test_room_exports_reject_mismatched_targets(saved,query):
    client,record=saved
    assert client.get(f"/api/generations/{record['id']}/export/pdf?{query}").status_code==404
