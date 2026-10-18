import csv
import io
from pathlib import Path
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
import pytest
from app.domain import Configuration, configuration_csv, parse_csv
from app.main import app

SAMPLE = Path(__file__).resolve().parents[1] / 'sample_data' / 'unisched_semesters_3_5_7.csv'

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('UNISCHED_DB', str(tmp_path / 'test.db'))
    from app import storage
    storage.initialize()
    storage.save_config(parse_csv(SAMPLE.read_bytes()).model_dump(), 0)
    with TestClient(app) as client:
        yield client

@pytest.fixture
def generated(client):
    response = client.post('/api/generations', json={'name':'Test semester', 'time_limit':20})
    assert response.status_code == 200, response.text
    record = response.json()
    assert record['result']['status'] == 'feasible', record
    return record

def test_startup_seeds_sample_once_and_serves_frontend(client):
    assert client.get('/api/health').json() == {'status':'ok','database':'sqlite'}
    data = client.get('/api/config').json()
    assert data['revision'] == 1
    assert len(data['configuration']['offerings']) == 54
    assert 'UniSched' in client.get('/').text
    assert client.get('/static/app.js').status_code == 200
    assert client.get('/api/sample.csv').content == SAMPLE.read_bytes()

def test_import_preview_does_not_mutate_database(client):
    before = client.get('/api/config').json()
    response = client.post('/api/import/preview', files={'file':('sample.csv',SAMPLE.read_bytes(),'text/csv')})
    assert response.status_code == 200
    assert response.json()['counts'] == {'teachers':18,'subjects':18,'rooms':15,'groups':9,'offerings':54,'departments':1,'batches':0,'elective_sets':0}
    assert response.json()['issues'] == []
    assert client.get('/api/config').json() == before
    result = client.put('/api/config', json={'revision':before['revision'],'configuration':response.json()['configuration']})
    assert result.status_code == 200
    assert result.json()['revision'] == 2

@pytest.mark.parametrize('content,expected',[
    (b'bad,headers\na,b\n','headers must match'),
    (SAMPLE.read_bytes().replace(b'09:00',b'25:00'),'24-hour time'),
    (SAMPLE.read_bytes().replace(b'false',b'maybe'),'Elective must be'),
    (SAMPLE.read_bytes().replace(b',18,',b',19,',1),'Conflicting repeated teachers'),
    (SAMPLE.read_bytes()+SAMPLE.read_bytes().splitlines(keepends=True)[1],'division/subject pair'),
    (b'\xff\xfe','UTF-8'),
    (b'x'*1_000_001,'smaller than 1 MB')])
def test_invalid_csv_is_rejected_without_changes(client,content,expected):
    before = client.get('/api/config').json()
    result = client.post('/api/import/preview', files={'file':('bad.csv',content)})
    assert result.status_code == 422
    assert expected in result.json()['detail']
    assert client.get('/api/config').json() == before

def test_csv_roundtrip():
    config = parse_csv(SAMPLE.read_bytes())
    exported = configuration_csv(config).encode()
    assert parse_csv(exported) == config

def test_concurrent_edits_and_invalid_references_are_rejected(client):
    data = client.get('/api/config').json()
    assert client.put('/api/config', json=data).status_code == 200
    assert client.put('/api/config', json=data).status_code == 409
    data = client.get('/api/config').json()
    data['configuration']['teachers'].pop()
    assert client.put('/api/config', json=data).status_code == 422
    assert len(client.get('/api/config').json()['configuration']['teachers']) == 18

def test_saved_snapshot_survives_configuration_change(client,generated):
    data = client.get('/api/config').json()
    data['configuration']['teachers'][0]['name'] = 'New Faculty Name'
    assert client.put('/api/config',json=data).status_code == 200
    snapshot = client.get(f"/api/generations/{generated['id']}").json()
    assert snapshot['config'] == generated['config']
    assert snapshot['config_revision'] != client.get('/api/config').json()['revision']
    assert len(client.get('/api/generations').json()) == 1

def test_division_pdf_and_excel_exports(client,generated):
    identifier = generated['id']
    xlsx = client.get(f'/api/generations/{identifier}/export/xlsx')
    assert xlsx.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert len(workbook.worksheets) == 9
    for sheet in workbook:
        assert 'Datta Meghe College of Engineering' in sheet['A1'].value
        assert sheet['B10'].value == 'RECESS'
        assert sheet.freeze_panes == 'B6'
        assert any('LAB' in str(cell.value) for row in sheet for cell in row)
    pdf = client.get(f'/api/generations/{identifier}/export/pdf')
    assert pdf.status_code == 200
    reader = PdfReader(io.BytesIO(pdf.content))
    assert len(reader.pages) == 9
    for page in reader.pages:
        text = page.extract_text()
        assert 'Datta Meghe College of Engineering' in text
        assert 'RECESS' in text

def test_professor_export_shows_only_their_classes(client,generated):
    teacher = generated['config']['teachers'][0]
    entries = [e for e in generated['result']['entries'] if e['teacher_id'] == teacher['id']]
    response = client.get(f"/api/generations/{generated['id']}/export/xlsx?view=teacher&target={teacher['id']}")
    workbook = load_workbook(io.BytesIO(response.content))
    assert len(workbook.worksheets) == 1
    sheet = workbook.active
    assert teacher['name'] in sheet['A2'].value
    cells = [cell.value for row in sheet.iter_rows(min_row=6,max_row=14,min_col=2,max_col=6) for cell in row]
    lessons = [text for text in cells if text and text != 'RECESS']
    assert len(lessons) == sum(e['duration'] for e in entries)
    assert any(text is None for text in cells)
    subjects = {s['id']:s for s in generated['config']['subjects']}
    valid_codes = {subjects[e['subject_id']]['code'] for e in entries}
    assert all(text.split(' · ')[0] in valid_codes for text in lessons)
    assert all('Div ' in text for text in lessons)
    response = client.get(f"/api/generations/{generated['id']}/export/pdf?view=teacher&target={teacher['id']}")
    pdf = PdfReader(io.BytesIO(response.content))
    assert len(pdf.pages) == 1
    assert teacher['name'] in pdf.pages[0].extract_text()

def test_failed_generation_cannot_be_exported(client):
    data = client.get('/api/config').json()
    data['configuration']['teachers'][0]['max_hours'] = 1
    client.put('/api/config',json=data)
    result = client.post('/api/generations',json={'name':'Impossible load'}).json()
    assert result['result']['status'] == 'infeasible'
    assert result['result']['issues'][0]['code'] == 'TEACHER_LOAD'
    assert client.get(f"/api/generations/{result['id']}/export/pdf").status_code == 409

def test_invalid_export_target(client,generated):
    assert client.get('/api/generations/999999').status_code == 404
    assert client.get(f"/api/generations/{generated['id']}/export/pdf?target=missing").status_code == 404
    assert client.get(f"/api/generations/{generated['id']}/export/pdf?view=unknown").status_code == 422

def test_untrusted_text_exports_as_text_not_formula(client,generated):
    from app.exports import excel, pdf
    generated['config']['teachers'][0]['name'] = '=1+1'
    generated['name'] = '<b>Untrusted & title</b>'
    body = excel(generated,'teacher',generated['config']['teachers'][0]['id'])
    workbook = load_workbook(io.BytesIO(body))
    assert all(c.data_type != 'f' for sheet in workbook for row in sheet for c in row)
    text = PdfReader(io.BytesIO(pdf(generated,'teacher',generated['config']['teachers'][0]['id']))).pages[0].extract_text()
    assert '<b>Untrusted & title</b>' in text

def test_export_unsafe_control_characters_are_rejected(client):
    data = client.get('/api/config').json()
    data['configuration']['teachers'][0]['name'] = 'Invalid\x01teacher'
    assert client.put('/api/config',json=data).status_code == 422
    assert client.post('/api/generations',json={'name':'Invalid\x01name'}).status_code == 422
